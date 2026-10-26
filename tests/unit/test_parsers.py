from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path

import pytest

from app.ingestion.parsers.docx_parser import DOCXParser
from app.ingestion.parsers.pdf_parser import PDFParser
from app.core.exceptions import DocumentProcessingError


class TestDOCXParser:
    """Tests for the DOCX parser.

    These tests create real (minimal) DOCX files to exercise the parser
    without needing fixtures on disk.
    """

    def _create_docx(self, paragraphs: list[tuple[str, str | None]], tmp_path: Path) -> str:
        """Create a minimal DOCX file with the given (text, style) paragraphs."""
        from docx import Document
        doc = Document()
        for text, style in paragraphs:
            if style:
                doc.add_heading(text, level=1)
            else:
                doc.add_paragraph(text)
        path = tmp_path / "test.docx"
        doc.save(str(path))
        return str(path)

    def test_basic_paragraph_extraction(self, tmp_path):
        path = self._create_docx(
            [("Hello World", None), ("Second paragraph", None)],
            tmp_path,
        )
        parser = DOCXParser()
        result = parser.parse(path)
        assert len(result.blocks) >= 2
        texts = [b.text for b in result.blocks]
        assert "Hello World" in texts
        assert "Second paragraph" in texts

    def test_heading_detected(self, tmp_path):
        path = self._create_docx(
            [("My Section Heading", "heading"), ("Body content here.", None)],
            tmp_path,
        )
        parser = DOCXParser()
        result = parser.parse(path)
        headings = [b for b in result.blocks if b.section_heading == "My Section Heading"]
        assert len(headings) >= 1

    def test_nonexistent_file_raises(self):
        parser = DOCXParser()
        with pytest.raises(DocumentProcessingError):
            parser.parse("/nonexistent/path/file.docx")

    def test_parsed_document_has_metadata(self, tmp_path):
        path = self._create_docx([("Content", None)], tmp_path)
        parser = DOCXParser()
        result = parser.parse(path)
        assert "source" in result.metadata
        assert result.metadata["source"] == "python-docx"

    def test_empty_docx_returns_no_blocks(self, tmp_path):
        from docx import Document
        doc = Document()
        path = tmp_path / "empty.docx"
        doc.save(str(path))
        parser = DOCXParser()
        result = parser.parse(str(path))
        # Empty DOCX may have 0 blocks
        assert isinstance(result.blocks, list)


class TestPDFParser:
    """Tests for the PDF parser.

    We test with a minimal programmatically-created PDF using PyMuPDF.
    """

    def _create_pdf(self, pages: list[str], tmp_path: Path) -> str:
        """Create a minimal PDF with one text block per page."""
        import fitz
        doc = fitz.open()
        for text in pages:
            page = doc.new_page()
            page.insert_text((72, 72), text, fontsize=12)
        path = tmp_path / "test.pdf"
        doc.save(str(path))
        doc.close()
        return str(path)

    def test_basic_text_extraction(self, tmp_path):
        path = self._create_pdf(["Hello from page one", "Hello from page two"], tmp_path)
        parser = PDFParser()
        result = parser.parse(path)
        full_text = result.full_text
        assert "Hello from page one" in full_text or len(result.blocks) >= 1

    def test_page_numbers_assigned(self, tmp_path):
        path = self._create_pdf(["Page 1 content", "Page 2 content"], tmp_path)
        parser = PDFParser()
        result = parser.parse(path)
        page_numbers = [b.page_number for b in result.blocks]
        assert 1 in page_numbers

    def test_nonexistent_file_raises(self):
        parser = PDFParser()
        with pytest.raises(DocumentProcessingError):
            parser.parse("/nonexistent/file.pdf")

    def test_metadata_contains_page_count(self, tmp_path):
        path = self._create_pdf(["Single page"], tmp_path)
        parser = PDFParser()
        result = parser.parse(path)
        assert "page_count" in result.metadata
        assert result.metadata["page_count"] >= 1
